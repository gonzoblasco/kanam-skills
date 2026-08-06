#!/usr/bin/env python3
"""export-novel.py — Export narrative content to DOCX, EPUB, or Markdown.

Usage:
  python3 export-novel.py <input-dir> [--format docx|epub|md] [--output <file>]

Takes a directory of chapter files (chapter-01.md, chapter-02.md, etc.)
and exports to the specified format.
"""

import sys
import os
import re
from pathlib import Path


def collect_chapters(input_dir: str) -> list[tuple[int, str, str]]:
    """Collect and sort chapter files from input directory."""
    chapters = []
    for f in Path(input_dir).iterdir():
        if f.suffix == ".md":
            # Extract chapter number from filename
            match = re.search(r"(\d+)", f.stem)
            num = int(match.group(1)) if match else 0
            chapters.append((num, f.stem, f.read_text()))
    chapters.sort(key=lambda x: x[0])
    return chapters


def export_markdown(chapters: list[tuple[int, str, str]], output: str):
    """Export as single merged markdown file."""
    with open(output, "w") as f:
        f.write("# Novel\n\n")
        for num, name, content in chapters:
            f.write(f"\n## {name}\n\n")
            f.write(content)
            f.write("\n\n---\n\n")
    print(f"✅ Exported to {output}")


def export_docx(chapters: list[tuple[int, str, str]], output: str):
    """Export as DOCX using python-docx if available."""
    try:
        from docx import Document
        from docx.shared import Pt, Inches
    except ImportError:
        print("❌ python-docx not installed. Install with: pip install python-docx")
        print("   Falling back to Markdown export.")
        export_markdown(chapters, output.replace(".docx", ".md"))
        return

    doc = Document()
    doc.add_heading("Novel", level=1)

    for num, name, content in chapters:
        doc.add_heading(name, level=2)
        for line in content.split("\n"):
            if line.strip():
                p = doc.add_paragraph(line.strip())
                p.style.font.size = Pt(12)
        doc.add_page_break()

    doc.save(output)
    print(f"✅ Exported to {output}")


def export_epub(chapters: list[tuple[int, str, str]], output: str):
    """Export as EPUB using pandoc if available."""
    import subprocess
    import tempfile

    # First create merged markdown
    with tempfile.NamedTemporaryFile(suffix=".md", mode="w", delete=False) as f:
        f.write("# Novel\n\n")
        for num, name, content in chapters:
            f.write(f"\n# {name}\n\n")
            f.write(content)
            f.write("\n\n")
        temp_md = f.name

    try:
        subprocess.run(
            ["pandoc", temp_md, "-o", output, "--from", "markdown", "--to", "epub"],
            check=True,
            capture_output=True,
        )
        print(f"✅ Exported to {output}")
    except FileNotFoundError:
        print("❌ pandoc not found. Install with: brew install pandoc")
        print("   Falling back to Markdown export.")
        export_markdown(chapters, output.replace(".epub", ".md"))
    finally:
        os.unlink(temp_md)


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    input_dir = sys.argv[1]
    output_format = "md"
    output_file = ""

    args = sys.argv[2:]
    while args:
        if args[0] == "--format" and len(args) > 1:
            output_format = args[1]
            args = args[2:]
        elif args[0] == "--output" and len(args) > 1:
            output_file = args[1]
            args = args[2:]
        else:
            args = args[1:]

    if not output_file:
        output_file = f"novel-{Path(input_dir).name}.{output_format}"

    chapters = collect_chapters(input_dir)
    if not chapters:
        print("❌ No chapter files found in input directory")
        sys.exit(1)

    print(f"📖 Found {len(chapters)} chapters")

    if output_format == "docx":
        export_docx(chapters, output_file)
    elif output_format == "epub":
        export_epub(chapters, output_file)
    else:
        export_markdown(chapters, output_file)


if __name__ == "__main__":
    main()
